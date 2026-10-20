#!/usr/bin/env python3
"""Check that nothing generated has drifted from the thing that generates it.

    python3 checks/generated.py

Five assertions:

  1. every generator reproduces the artefacts committed beside it, byte for byte
  2. no committed artefact is one its generator no longer writes
  3. the service design boards re-render to the committed SVG
  4. the scan reached every generator, so a passing run means something
  5. the "In the playbook" row of every journey sheet in the Miro workbook reads as the board does

This file exists because a generator was eight days ahead of its artefact and nothing saw it. A
diagram generator had been edited to call NHSBSA's sheet by its proper name, the Analytical
Assurance Decision Tool, and the `.drawio` it wrote was never rebuilt, so the picture the playbook
shipped still said "the assurance level decision sheet". The whole suite passed. It had to, because
every check downstream was consistent with the stale artefact: one asserted that each `.drawio`
re-renders to the `.svg` beside it, and it did; another asserted the labels fit their boxes, and
they did; the alt text matched the drawn title, because both were stale together. A chain of checks
that each compare one step to the next is blind to the step before the first one.

So assertion 1 is the missing link in that chain: the generator is the source of truth, and this
check asserts the `.drawio` matches it.

Assertion 2 is the same drift in the other direction, and it is cheap once the first is written.
A board renamed in its generator leaves the old file on disk, where it keeps passing every other
check while belonging to nothing.

Assertion 3 covers the rendering step: the service design boards live under `artefacts/` and are
shipped in the workshop pack, and they are drawn by the same renderer as anything else, so the same
staleness is available to them. The playbook's own diagrams are written by `build/diagrams.py`
and compared here as sources, and rendered and compared by `checks/diagrams.py`.

Assertion 4 is here because the first three are all of the form "run the generator and compare".
If a generator stops being importable, or its `main()` loses the `out` parameter these call it
with, every comparison quietly tests nothing and the check still prints ok. It fails instead.

Both generators are deterministic by construction: `build/drawio.py` writes ids from a counter
rather than from randomness, and neither generator puts a date in what it draws. The stamp along
the foot of a service design board names its author, the organisation and the Statement of Works and nothing that
changes between two runs. That is what makes a byte comparison the right test here, and it is
worth saying out loud, because the moment either generator stamps a build date this check starts
failing every day and the fix is to exclude the stamp rather than to delete the check.

The fifth is here because the workbook cannot be compared byte for byte: it carries the time it
was written, so two runs differ. It can be compared in the one respect that went wrong. A label in
the "In the playbook" row is a link label, and the boards break it over two lines only to shape the
box (`Deciding your<br>assurance level`). The workbook turned every break into a colon, so a screen
reader said "Deciding your colon assurance level" in six cells, and "The Analysis Senior colon
Responsible Officer" in a seventh that the lowercase test below cannot see. A colon belongs only
where the break separates two things, which in these labels means after a head: "Phase 3", "R10",
"The commissioner", "The critical friend". So a cell fails if it has a colon and a lowercase word
after it that the board's own text does not have, or a colon that follows anything but a head.

PNGs are deliberately not compared: they are
rasterised, and two rasterisers differ on anti-aliasing by a percent or two of pixels without
either being wrong. The `.drawio` and the `.svg` are exact artefacts, so those are what is
compared.

This needs `build/`, which the release package does not carry, so it skips there and says so
rather than passing quietly. The fifth assertion also needs openpyxl, which reads the workbook.
"""
import contextlib
import importlib.util
import io
import pathlib
import re
import sys
import tempfile

ROOT = pathlib.Path(__file__).resolve().parent.parent

# Each entry is: the generator, the directory it writes into, and what it writes there. The
# generator's main() must take an `out` directory, which is the whole of what this check needs
# from it.
GENERATORS = [
    ('build/diagrams.py', 'static/diagrams', '*.drawio'),
    ('build/build-service-design.py', 'artefacts/service-design', '*.drawio'),
]

# Rendering steps this check owns. The playbook's own diagrams are rendered and compared by
# checks/diagrams.py, so they are not repeated here.
RENDERED = [('artefacts/service-design', 'artefacts/service-design/png')]

# Assertion 5. The workbook, and what may stand before a colon in one of its link labels.
WORKBOOK = 'artefacts/service-design/miro-paste-grids.xlsx'
LABEL_HEADS = re.compile(r'Phase \d+|R\d+|The commissioner|The critical friend')

fail, skip, done = [], [], []


def load(rel):
    """Import a build script by path. Several have dashes in the name, so none of them import."""
    path = ROOT / rel
    spec = importlib.util.spec_from_file_location('gen_' + path.stem.replace('-', '_'), path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def compare(fresh_dir, live_dir, pattern, what):
    """Compare what a run just wrote against what is committed, in both directions."""
    fresh = {f.name: f for f in sorted(fresh_dir.glob(pattern))}
    live = {f.name: f for f in sorted(live_dir.glob(pattern))}
    if not fresh:
        fail.append('%s wrote no %s, so nothing was compared' % (what, pattern))
        return 0

    for name in sorted(fresh):                                                # 1
        if name not in live:
            fail.append('%s writes %s and there is no such file in %s. The generator is ahead of '
                        'the artefact: run it and commit what it writes'
                        % (what, name, live_dir.relative_to(ROOT)))
        elif fresh[name].read_bytes() != live[name].read_bytes():
            fail.append('%s does not match what %s writes now. The generator has been edited and '
                        'the artefact beside it has not been rebuilt, so the shipped version is '
                        'the older one' % (live_dir.relative_to(ROOT) / name, what))

    for name in sorted(live):                                                 # 2
        if name not in fresh:
            fail.append('%s is committed and %s no longer writes it. Either the generator lost a '
                        'diagram or the file is a leftover from one that was renamed'
                        % (live_dir.relative_to(ROOT) / name, what))
    return len(fresh)


def check_playbook_row(sd):
    """Assertion 5. Returns how many cells it read."""
    from openpyxl import load_workbook
    path = ROOT / WORKBOOK
    if not path.exists():
        fail.append('%s is missing' % WORKBOOK)
        return 0
    sheets = {ws['A1'].value: ws for ws in load_workbook(path).worksheets}
    read = 0
    for spec in sd.JOURNEYS:
        ws = sheets.get('%s: journey through the playbook' % spec['role'])
        row = next((r for r in range(1, (ws.max_row if ws else 0) + 1)
                    if ws.cell(r, 1).value == 'In the playbook'), None)
        if row is None:
            fail.append('%s has no "In the playbook" row for %s' % (WORKBOOK, spec['role']))
            continue
        for n, stop in enumerate(spec['stops'], start=1):
            cell = ws.cell(row, n + 1).value or ''
            board = sd.words(stop[1])
            read += 1
            where = '%s, %s, stop %d' % (WORKBOOK, spec['role'], n)
            for m in re.finditer(r':\s+[a-z][\w-]*', cell):
                if m.group(0) not in board:
                    fail.append('%s reads %r and the board reads %r. The break is a line wrap, and '
                                'a screen reader says "colon" at it' % (where, cell, board))
                    break
            else:
                head = cell.split(':')[0] if ':' in cell else None
                if head is not None and not LABEL_HEADS.fullmatch(head):
                    fail.append('%s reads %r and the colon follows %r, which is not a head such as '
                                '"Phase 3". Where a break is only a line wrap it is a space'
                                % (where, cell, head))
    return read


def main():
    if not (ROOT / 'build').is_dir():
        print('  --  nothing was re-generated: build/ is not in the release package, so the '
              'generators are not here')
        return

    with tempfile.TemporaryDirectory() as tmp:
        for rel, out_rel, pattern in GENERATORS:
            live = ROOT / out_rel
            if not live.is_dir():
                fail.append('%s writes into %s, which does not exist' % (rel, out_rel))
                continue
            fresh = pathlib.Path(tmp) / pathlib.Path(rel).stem
            try:
                module = load(rel)
                # The generators print a line per artefact. That is right when a person runs them
                # and noise when this does.
                with contextlib.redirect_stdout(io.StringIO()):
                    module.main(out=fresh)
            except TypeError as e:
                fail.append("%s could not be run into a temporary directory: %s. Its main() has "
                            "to take an `out` directory, or nothing here compares anything" % (rel, e))
                continue
            except Exception as e:                        # noqa: BLE001 - any failure is a failure
                fail.append('%s did not run: %s: %s' % (rel, type(e).__name__, e))
                continue
            n = compare(fresh, live, pattern, rel)
            if n:
                done.append((rel, out_rel, n))

        for src_rel, out_rel in RENDERED:                                     # 3
            src, live = ROOT / src_rel, ROOT / out_rel
            sources = sorted(src.glob('*.drawio'))
            if not sources:
                skip.append('no .drawio in %s to re-render' % src_rel)
                continue
            fresh = pathlib.Path(tmp) / ('render-' + pathlib.Path(out_rel).name)
            fresh.mkdir(parents=True, exist_ok=True)
            try:
                rd = load('build/render-diagrams.py')
                with contextlib.redirect_stdout(io.StringIO()):
                    rd.main(src=src, out=fresh)
            except Exception as e:                        # noqa: BLE001 - any failure is a failure
                fail.append('the boards in %s did not re-render: %s: %s'
                            % (src_rel, type(e).__name__, e))
                continue
            for s in sources:
                new, old = fresh / (s.stem + '.svg'), live / (s.stem + '.svg')
                if not old.exists():
                    fail.append('%s has no .svg in %s' % (s.name, out_rel))
                elif not new.exists():
                    fail.append('%s did not render' % s.name)
                elif new.read_bytes() != old.read_bytes():
                    fail.append('%s does not match its source. The .drawio has been edited and '
                                'the .svg beside it has not been rebuilt'
                                % (pathlib.Path(out_rel) / (s.stem + '.svg')))
            done.append(('build/render-diagrams.py', out_rel, len(sources)))

    try:                                                                      # 5
        cells = check_playbook_row(load('build/build-service-design.py'))
    except Exception as e:                            # noqa: BLE001 - any failure is a failure
        cells = 0
        fail.append('the workbook could not be read against the boards: %s: %s'
                    % (type(e).__name__, e))

    for s in skip:
        print('  --  %s' % s)
    if fail:
        print('generated FAILED')
        for f in fail:
            print('  X   %s' % f)
        sys.exit(1)

    # 4. Every comparison above is "run it and diff it", and a generator that cannot be run
    # compares nothing while still printing ok. Nothing below is reached unless each one ran.
    ran = {rel for rel, _, _ in done}
    missing = [rel for rel, _, _ in GENERATORS if rel not in ran]
    if missing:
        raise SystemExit('generated FAILED\n  X   %d generator(s) produced nothing to compare, so '
                         'this check tested nothing: %s' % (len(missing), ', '.join(missing)))

    for rel, out_rel, n in done:
        print('  ok  %-32s %2d artefact(s) in %s, every one matching' % (rel, n, out_rel))
    print('  ok  no artefact in those directories that its generator no longer writes')
    print('  ok  %s: %d "In the playbook" cells, none with a colon that is only a line wrap'
          % (pathlib.Path(WORKBOOK).name, cells))


if __name__ == '__main__':
    main()
